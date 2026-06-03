#!/usr/bin/env python3
"""Generate a single-page PPTX for the Demo slide (real numbers, white theme).
Standalone — does not touch the main deck. Output: demo_onepage.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x66, 0x66, 0x66)
CARD = RGBColor(0xF2, 0xF2, 0xF2)
BORDER = RGBColor(0xD9, 0xD9, 0xD9)
BLUE = RGBColor(0x2E, 0x86, 0xAB)
CORAL = RGBColor(0xE0, 0x7A, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])   # blank

# white background
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = WHITE; bg.line.fill.background()
bg.shadow.inherit = False
s.shapes._spTree.remove(bg._element); s.shapes._spTree.insert(2, bg._element)


def textbox(x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(0); tf.margin_right = Pt(0)
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    return tf


def run(p, text, size, bold=False, color=INK, font="Calibri"):
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = font
    return r


# ---- title ----
tf = textbox(Inches(0.6), Inches(0.35), Inches(8), Inches(0.9))
run(tf.paragraphs[0], "Demo", 40, bold=True)
# underline rule
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.25),
                          Inches(12.1), Pt(1.5))
line.fill.solid(); line.fill.fore_color.rgb = BORDER; line.line.fill.background()
line.shadow.inherit = False


def card(x, w):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.7),
                           w, Inches(4.9))
    c.fill.solid(); c.fill.fore_color.rgb = CARD
    c.line.color.rgb = BORDER; c.line.width = Pt(1)
    c.shadow.inherit = False
    try:
        c.adjustments[0] = 0.04
    except Exception:
        pass
    return c


# ---- left card: steps ----
card(Inches(0.6), Inches(5.75))
tf = textbox(Inches(1.0), Inches(2.05), Inches(5.0), Inches(0.6))
run(tf.paragraphs[0], "Demo A: tool-pause resume", 21, bold=True)

# (kind, text): "step" = numbered; coral/blue/gray = indented sub-line
steps = [
    ("step",  "Start a long-context request, stream output"),
    ("step",  "Trigger a 10 s tool / API pause"),
    ("gray",  "GPU KV released — usage drops to 0"),
    ("step",  "Resume two ways (same engine):"),
    ("coral", "recompute — re-prefill the whole context"),
    ("blue",  "Ferry — reload exact KV from host (1 step)"),
    ("step",  "Both continue seamlessly — even mid-word"),
    ("step",  "Compare resume cost — Ferry ~15x faster"),
]
SUBCOL = {"gray": GRAY, "coral": CORAL, "blue": BLUE}
tf = textbox(Inches(1.0), Inches(2.9), Inches(5.2), Inches(3.6))
n = 0
for i, (kind, txt) in enumerate(steps):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(7)
    if kind == "step":
        n += 1
        run(p, f"{n}.  {txt}", 15.5)
    else:
        run(p, "      " + txt, 14, color=SUBCOL[kind])

# ---- right card: terminal output (real numbers) ----
card(Inches(6.95), Inches(5.75))
tf = textbox(Inches(7.3), Inches(2.05), Inches(5.0), Inches(0.6))
run(tf.paragraphs[0], "terminal output", 19, bold=True, font="Consolas")

MONO = "Consolas"
tf = textbox(Inches(7.3), Inches(2.95), Inches(5.2), Inches(3.4))


def mono(p, parts):
    """parts: list of (text, color) tuples."""
    for txt, col in parts:
        run(p, txt, 13.5, color=col, font=MONO)


rows = [
    [("GPU KV cache usage  (the pause)", INK)],
    [("  running  14.7%", INK), ("   KV on GPU", GRAY)],
    [("  PAUSE     0.0%", CORAL), ("   freed -> host RAM", GRAY)],
    [("  reload   14.8%", BLUE), ("   back from host", GRAY)],
    [("", INK)],
    [("resume cost  (time-to-first-token)", INK)],
    [("  recompute  re-prefill ~9k tok", CORAL), ("  2.62 s", INK)],
    [("  Ferry      reload, 1 step", BLUE), ("     0.17 s", INK)],
    [("                            ", INK), ("    ~15x faster", GRAY)],
]
for i, parts in enumerate(rows):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(4)
    mono(p, parts)

out = Path(__file__).resolve().parent / "demo_onepage.pptx"
prs.save(out)
print("saved:", out)
