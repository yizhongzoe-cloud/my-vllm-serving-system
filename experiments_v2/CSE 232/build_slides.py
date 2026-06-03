#!/usr/bin/env python3
"""Build the Ferry CSE 232 project slide deck (simple white/black, bullets).
Speaker notes carry a Chinese summary + an English speaker script per slide.
Figures are reused from experiments_v2/figures/. Run from repo root or here.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.normpath(os.path.join(HERE, "..", "figures"))

INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x2E, 0x86, 0xAB)
CORAL = RGBColor(0xE0, 0x7A, 0x5F)
GRAY = RGBColor(0x99, 0x99, 0x99)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
DARKPANEL = RGBColor(0x26, 0x33, 0x38)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide():
    return prs.slides.add_slide(BLANK)


def notes(s, zh, en):
    tf = s.notes_slide.notes_text_frame
    tf.text = "【中文】" + zh + "\n\n【English script】" + en


def title(s, text, accent=BLUE):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = INK
    r.font.name = "Arial"
    # accent underline bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.28),
                             Inches(2.2), Pt(4))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False
    return tb


def _set(run, size, color=INK, bold=False, name="Arial"):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.name = name


def bullets(s, items, left=0.7, top=1.7, width=12.0, height=5.3, size=19,
            gap=8):
    """items: list of (level, text) or str. level 0/1/2. text may start with
    '*' to render bold accent (a highlight line)."""
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width),
                              Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        if isinstance(it, tuple):
            lvl, txt = it
        else:
            lvl, txt = 0, it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        p.space_after = Pt(gap)
        hl = txt.startswith("*")
        if hl:
            txt = txt[1:]
        marker = "" if lvl == 0 else ("–  " if lvl == 1 else "·  ")
        bullet = ("●  " if lvl == 0 else marker)
        r = p.add_run(); r.text = bullet + txt
        _set(r, size - (lvl * 1), CORAL if hl else INK, bold=(lvl == 0 or hl))
    return tb


def codebox(s, lines, left=0.7, top=1.8, width=7.2, height=4.6, title_txt=None):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left),
                             Inches(top), Inches(width), Inches(height))
    box.fill.solid(); box.fill.fore_color.rgb = DARKPANEL
    box.line.fill.background(); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.18)
    tf.margin_right = Inches(0.2)
    first = True
    if title_txt:
        p = tf.paragraphs[0]; first = False
        r = p.add_run(); r.text = title_txt
        _set(r, 15, RGBColor(0xD9, 0xB3, 0x8C), bold=True, name="Consolas")
        p.space_after = Pt(8)
    for fn, desc in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        r = p.add_run(); r.text = fn
        _set(r, 14.5, RGBColor(0x8E, 0xC9, 0xE6), bold=True, name="Consolas")
        if desc:
            r2 = p.add_run(); r2.text = "  " + desc
            _set(r2, 13.5, RGBColor(0xE8, 0xE8, 0xE8), name="Consolas")
    return box


def card(s, left, top, width, height, head, body, head_color=BLUE):
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                             Inches(width), Inches(height))
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT
    box.line.fill.background(); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.18); tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = head
    _set(r, 14, head_color, bold=True); p.space_after = Pt(4)
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = body
    _set(r2, 12.5, INK)
    return box


def picture(s, name, left, top, width=None, height=None):
    kw = {}
    if width: kw["width"] = Inches(width)
    if height: kw["height"] = Inches(height)
    return s.shapes.add_picture(os.path.join(FIG, name), Inches(left),
                                Inches(top), **kw)


def caption(s, text, left=0.7, top=6.9, width=12.0, size=12, color=GRAY):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width),
                              Inches(0.5))
    p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = text
    _set(r, size, color)
    return tb


def box_with_text(s, left, top, w, h, lines, fill=None, line_color=INK,
                  align=PP_ALIGN.CENTER, size=12, bold0=True):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                            Inches(w), Inches(h))
    if fill is None:
        sp.fill.solid(); sp.fill.fore_color.rgb = WHITE
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line_color; sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        _set(r, size, INK, bold=(bold0 and i == 0))
    return sp


def arrow(s, x1, y1, x2, y2, color=BLUE, label=None, lx=None, ly=None,
          width=2.0):
    cn = s.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2),
                                Inches(y2))
    cn.line.color.rgb = color; cn.line.width = Pt(width)
    cn.shadow.inherit = False
    le = cn.line._get_or_add_ln()
    tail = le.makeelement(qn('a:tailEnd'),
                          {'type': 'triangle', 'w': 'med', 'len': 'med'})
    le.append(tail)
    if label:
        tb = s.shapes.add_textbox(Inches(lx if lx is not None else x1),
                                  Inches(ly if ly is not None else y1),
                                  Inches(2.6), Inches(0.35))
        p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label; _set(r, 11, color, bold=True)
    return cn


# ============================== SLIDE 1: TITLE ==============================
s = slide()
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.0))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run()
r.text = "Ferry"
_set(r, 54, BLUE, bold=True)
p2 = tf.add_paragraph(); r2 = p2.add_run()
r2.text = "Decoupling KV-State Lifetime from GPU Residency\nin Long-Context LLM Serving"
_set(r2, 26, INK, bold=True)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(4.35),
                         Inches(3.0), Pt(4))
bar.fill.solid(); bar.fill.fore_color.rgb = CORAL; bar.line.fill.background()
bar.shadow.inherit = False
tb2 = s.shapes.add_textbox(Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.5))
tf2 = tb2.text_frame
for i, ln in enumerate(["Yi Zhong   ·   Chen Yang   ·   Guanghua Sun",
                        "CSE 232 — AI Systems Project",
                        "Prototype on vLLM · Qwen2.5 · NVIDIA A6000"]):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    r = p.add_run(); r.text = ln
    _set(r, 16 if i == 0 else 14, INK if i == 0 else GRAY, bold=(i == 0))
notes(s,
      "标题页。一句话点题:Ferry 把「KV 的生命周期」和「它待在 GPU 上」这两件事解耦。作者/课程/平台占位——名字和单位你确认一下(论文里是 Yi Zhong / Chen Yang / Guanghua Sun)。",
      "Hi everyone. Our project is Ferry, a serving-system mechanism for long-context LLMs. The one-line idea is in the title: we decouple the lifetime of a request's KV cache from whether it currently sits in GPU memory. I'll motivate why that matters, walk through the design and implementation, then show results and a short demo.")

# ============================== SLIDE 2: DILEMMA ==============================
s = slide(); title(s, "The KV-Residency Dilemma")
bullets(s, [
    "Long-context requests carry a huge KV cache — tens of GB for a single request.",
    "Modern requests are not pure generation: they pause mid-execution for retrieval, tool calls, and APIs (interception).",
    "When a request pauses, its KV cache goes idle but is still needed to continue.",
    "*The dilemma — what do you do with idle KV during the wait?",
    (1, "Keep it on GPU → fast resume, but ties up scarce HBM for the whole pause."),
    (1, "Discard it → frees memory, but forces a full re-prefill of the long prompt on resume."),
    "This is a memory + scheduling problem in AI serving systems — not a model-quality problem.",
], size=19)
notes(s,
      "动机页。讲清楚问题:长上下文请求 KV 很大;它们会因为工具/检索/API 而中途暂停;暂停时这块 KV 闲着但又不能丢。于是两难——留在显存(快但占内存)还是丢掉(省内存但恢复要重算)。强调这是系统层面的内存/调度问题,契合课程主题。",
      "Long-context requests carry enormous KV caches. And increasingly they pause in the middle — to call a tool, hit a retrieval system, or wait on an API. During that pause the KV cache sits idle but you still need it to continue. That creates a dilemma: keep it on the GPU and you waste scarce memory for the whole wait; drop it and you pay a full re-prefill when the request resumes. This is fundamentally a memory-management and scheduling problem, which is squarely an AI-systems problem.")

# ============================== SLIDE 3: COST ==============================
s = slide(); title(s, "Why It Matters: the Cost of Long-Context State")
bullets(s, [
    "Prefill dominates long-context cost. On Qwen2.5-7B / A6000:",
    (1, "4 concurrent 16K-token requests ≈ 10 s of prefill."),
    (1, "4 concurrent 64K-token requests ≈ 70+ s of prefill."),
    "*Discard-and-recompute repeats ALL of that work on every resume.",
    "Tool-induced idle gaps span 50 ms (local code) to 30+ s (web / API calls).",
    "*Reactive swap (copy KV to host AT the pause) avoids recompute — but the GPU→host transfer sits on the critical path: can't free until swap-out finishes, can't resume until swap-in finishes.",
    "Every existing option reacts AFTER the pause. None prepares in advance.",
], size=18)
notes(s,
      "继续动机,给「贵」一个具体数字。长上下文 prefill 极贵(4×16K≈10s,4×64K≈70s)。丢了就得全重算。暂停时长从 50ms 到 30s 不等。反应式 swap 虽然不重算,但把传输放在关键路径上。核心痛点:现有做法都是「暂停之后才反应」,没人提前准备。",
      "Why does this matter quantitatively? Prefill dominates at long context — on a 7B model, four 16K requests take about ten seconds to prefill, and four 64K requests take over seventy. If you discard the KV, you repeat all of that on resume. Tool pauses themselves range from milliseconds to tens of seconds. Reactive swapping avoids the recompute, but it copies the KV out only after the pause begins, so the transfer is on the critical path. The common flaw: every existing strategy reacts after the pause — none prepares in advance.")

# ============================== SLIDE 4: RESEARCH QUESTIONS ==============================
s = slide(); title(s, "Research Questions")
bullets(s, [
    "*Q1 — Mechanism cost.  What is the steady-state overhead of continuously checkpointing KV during normal decode? How does reload compare to full re-prefill in isolation?",
    "",
    "*Q2 — Interception.  On intercepted long-context workloads, how does Ferry compare to GPU-resident retention, discard-and-recompute, and reactive swap — in resume latency and goodput?",
    "",
    "*Q3 — Generality.  Does the same substrate enable recovery from other events (e.g. worker failure) without paying a full re-prefill?",
], size=21)
notes(s,
      "把全文收敛成三个问题:Q1 机制本身贵不贵(checkpoint 开销 + reload vs re-prefill);Q2 拦截场景下对比三种基线;Q3 同一套底座能不能顺带解决故障恢复。后面 evaluation 就按这三问回答。",
      "We frame the work around three questions. First, mechanism cost: is continuous checkpointing cheap enough, and how much cheaper is reload than re-prefill? Second, interception: on pause-heavy long-context workloads, how does Ferry compare against keeping KV resident, discarding and recomputing, and reactive swap? And third, generality: does the very same substrate also let us recover from a worker failure without re-prefilling? Our evaluation answers these in order.")

# ============================== SLIDE 5: RELATED / EXISTING ==============================
s = slide(); title(s, "Existing Reactive Strategies — and Why They Fall Short")
rows = [
    ("Retain on GPU", "Instant resume", "Holds scarce HBM for the entire pause"),
    ("Discard + recompute", "Frees memory", "Full re-prefill on resume (10–70 s at long ctx)"),
    ("Reactive swap (INFERCEPT)", "Avoids recompute", "GPU→host copy on the critical path of the pause"),
]
tbl = s.shapes.add_table(4, 3, Inches(0.7), Inches(1.8), Inches(12.0),
                         Inches(2.8)).table
for j, h in enumerate(["Strategy", "Upside", "Downside"]):
    c = tbl.cell(0, j); c.text = h
    c.text_frame.paragraphs[0].runs[0].font.bold = True
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    c.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
    c.fill.solid(); c.fill.fore_color.rgb = BLUE
for i, row in enumerate(rows, start=1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j); c.text = val
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(14); r.font.color.rgb = INK
        if j == 0: r.font.bold = True
        c.fill.solid(); c.fill.fore_color.rgb = WHITE
bullets(s, [
    "*All three treat KV movement as a reaction to the pause — the work happens on the critical path, after the request has already blocked.",
    "Paged KV & prefix caching make residency efficient, but do not make KV state independent of the GPU pool.",
], top=4.9, size=18)
notes(s,
      "相关工作/现状,一张表说清三种反应式策略各自的优劣:留显存、丢了重算、反应式 swap。共性缺陷:都在关键路径上、暂停之后才动。paged KV / prefix caching 只是让「在显存里」更高效,并没有让 KV 脱离 GPU。这给 Ferry 留出位置。",
      "Here are the three strategies systems use today. Retain on GPU gives instant resume but wastes memory. Discard and recompute frees memory but pays a full re-prefill. Reactive swap, like INFERCEPT, avoids recompute but the copy is on the critical path. The common weakness is in the last column — all of them act after the pause, on the critical path. Paged attention and prefix caching make on-GPU management efficient, but they never make a request's KV independent of the GPU pool. That's the gap Ferry fills.")

# ============================== SLIDE 6: KEY INSIGHT ==============================
s = slide(); title(s, "Key Insight: KV Is Append-Only at Block Granularity")
bullets(s, [
    "A running request builds its KV incrementally; at vLLM's block granularity (16 tokens), a completed block never changes — later decoding appends new blocks.",
    "*So we don't have to wait for the pause to preserve state — completed blocks can be copied to host memory in the background, while the request is still decoding.",
    "Decode is HBM-bound and leaves PCIe largely idle → spare bandwidth to publish blocks for free.",
    "When the request later pauses, most of its KV is already on the host: suspension just records a manifest entry and releases the GPU blocks.",
    "*Analogy: VM pre-copy live migration — move state before the switch is needed.",
], size=18)
notes(s,
      "系统设计的核心 observation。KV 是 append-only 的:块写完就不变。所以不必等暂停才保存,可以边解码边在后台把完成的块拷到 host。decode 是 HBM-bound,PCIe 基本闲着,正好用这条带宽。等真暂停时,大部分 KV 已在 host,暂停只需记 manifest + 释放显存。类比 VM 预拷贝热迁移。",
      "Here is the key observation that makes Ferry work. A request's KV is append-only: once a 16-token block is finished, it never changes — decoding only appends new blocks. That means we don't need to wait for the pause to save state. We can copy completed blocks to host memory in the background while the request is still decoding. And decoding is memory-bandwidth bound, so the PCIe link is mostly idle — we publish blocks essentially for free. By the time the request pauses, most of its KV is already on the host, so suspension just records a manifest entry and releases the GPU blocks. It's the same idea as pre-copy live migration of VMs: move the state before you need to switch.")

# ============================== SLIDE 7: ARCHITECTURE ==============================
s = slide(); title(s, "Architecture: a Checkpoint–Reload Substrate on vLLM")
# two engine boxes
box_with_text(s, 0.8, 1.7, 3.6, 1.7,
              ["Engine 1", "Scheduler · Model Executor",
               "GPU: weights + KV blocks"], size=13)
box_with_text(s, 8.9, 1.7, 3.6, 1.7,
              ["Engine 2", "Scheduler · Model Executor",
               "GPU: weights + KV blocks"], size=13)
# host store
box_with_text(s, 2.6, 5.1, 8.1, 1.3,
              ["Host-Side Checkpoint Store",
               "completed KV blocks  +  manifest (gen · offset · block map)"],
              fill=LIGHT, size=14)
# arrows
arrow(s, 2.4, 3.4, 4.3, 5.1, color=BLUE, label="continuous\ncheckpoint",
      lx=1.0, ly=4.0)
arrow(s, 8.9, 5.1, 10.8, 3.4, color=CORAL, label="reload\n(same / other engine)",
      lx=10.3, ly=4.0)
# pause marker
box_with_text(s, 5.4, 2.0, 2.5, 0.9,
              ["⏸ external call", "release GPU KV"], fill=WHITE,
              line_color=GRAY, size=12)
caption(s, "Continuously publish completed KV blocks → on pause release GPU KV → on resume reload from host + replay bounded suffix.  Three parts: checkpoint store · reload path · suspend–resume path.")
notes(s,
      "总架构图。两台引擎 + 一个 host 端 checkpoint store。运行时引擎持续把完成的块发布到 store(蓝箭头);暂停时释放 GPU 上的 KV;恢复时从 store 把块 reload 回来(珊瑚箭头),同一台或另一台引擎都行。底座三部分:checkpoint store、reload 路径、suspend-resume 路径。这是系统设计部分的主图,要多讲。",
      "This is the overall architecture — a checkpoint-reload substrate sitting on top of vLLM. We run two engines, each with its scheduler, executor, and GPU memory. The new piece is the host-side checkpoint store at the bottom. While a request runs, its engine continuously publishes completed KV blocks to that store — the blue arrow. When the request hits an external call, we release its GPU KV blocks. When it resumes, we reload the blocks from the host store — the coral arrow — onto the same engine or a different one. The substrate has three parts: the checkpoint store, the reload path, and the suspend-resume path. Because reload doesn't care why the blocks were released, the same machinery also handles reroute and failure, which I'll come back to.")

# ============================== SLIDE 8: CHECKPOINT STORE ==============================
s = slide(); title(s, "Component 1 — Host-Side Checkpoint Store")
bullets(s, [
    "Accumulates completed KV blocks in host RAM ahead of any pause, moving GPU→host transfer off the suspension path.",
    "Block granularity (16 tokens) = vLLM's KV allocation unit = a natural atomic publish unit.",
    "*A block is published only after BOTH the GPU→host copy AND the manifest update finish — reload never sees a partial block.",
    "Manifest tracks: latest generation · covered-token offset · logical→host block mapping (to rebuild the paged layout on reload).",
    "Why host RAM: far larger than HBM, and reloads are far cheaper than re-prefill — vs peer-GPU (NVLink + extra HBM) or SSD (latency back on the critical path).",
], size=18)
notes(s,
      "组件一:host 端 checkpoint store。提前把完成块攒在 host RAM,把 GPU→host 传输挪出暂停路径。以块(16 token)为原子发布单位。关键:块只有在「拷贝完成 + manifest 更新完成」后才算发布,保证 reload 不会读到半个块。manifest 记三样:代号、覆盖到第几个 token、逻辑块到 host 块的映射。选 host RAM 而不是 peer-GPU/SSD 的理由也讲一下。",
      "Component one is the host-side checkpoint store. It accumulates completed KV blocks in host RAM before any pause, which is what takes the transfer off the suspension path. We checkpoint at vLLM's block granularity — 16 tokens — because that's already the allocation unit and gives us a clean atomic publish unit. A block counts as published only after both the copy and the manifest update finish, so reload never reads a half-copied block. The manifest tracks the latest generation, how many tokens are covered, and the mapping from logical blocks to host blocks so we can rebuild the paged layout. We use host RAM because it's much larger than HBM and reloads from it are far cheaper than re-prefill — peer-GPU staging burns NVLink and target HBM, and SSD would put latency back on the critical path.")

# ============================== SLIDE 9: RELOAD + REPLAY ==============================
s = slide(); title(s, "Component 2 — Checkpoint Reload + Bounded Replay")
bullets(s, [
    "On resume, restore from the latest complete checkpoint instead of re-prefilling the prompt:",
    (1, "Read manifest → allocate fresh GPU KV blocks → copy K/V back → rebuild paged-attention layout."),
    "*Because checkpoints are published at block boundaries, a resumed request lacks at most the tokens after the last published block.",
    "→ Replay becomes a bounded suffix repair, not a second prefill.",
    "*Resume cost ≈ host→GPU reload  +  small bounded replay  —  NOT reprocessing the long prompt.",
], size=19)
notes(s,
      "组件二:reload + 有界 replay。恢复时不重算 prompt,而是读 manifest、分配新块、把 K/V 拷回、重建 paged layout。因为只在块边界发布,最多缺最后一个块之后的几个 token,所以 replay 是「有界的尾巴修补」,不是第二次 prefill。恢复代价≈ reload + 小 replay,跟长 prompt 无关。",
      "Component two is reload with bounded replay. When a request resumes, we restore it from its latest checkpoint rather than re-prefilling. We read the manifest, allocate fresh GPU blocks, copy the K and V tensors back, and rebuild the paged-attention layout. Because we only publish at block boundaries, a resumed request is missing at most the handful of tokens generated after its last published block. So replay is a short, bounded suffix repair — not a second prefill. The resume cost is therefore the host-to-GPU reload plus a tiny replay, completely independent of how long the original prompt was.")

# ============================== SLIDE 10: SUSPEND-RESUME ==============================
s = slide(); title(s, "Component 3 — Suspend / Resume for Interception")
bullets(s, [
    "When an intercepted request issues an external call, its KV becomes idle until the result returns.",
    "*Suspend decision (a tradeoff): suspend only if a checkpoint exists AND the idle GPU blocks would block waiting work; otherwise retain.",
    (1, "Short waits → retain (avoid reload overhead)."),
    (1, "Long waits → suspend (freed memory can serve other requests)."),
    "Suspend = record external-call state + checkpoint generation, drop from running set, release GPU KV. Only lightweight metadata remains.",
    "On result return, the request re-enters the scheduler as resumable work and reloads once KV space is free.",
], size=18)
notes(s,
      "组件三:拦截请求的 suspend/resume,以及那个权衡。外部调用时 KV 闲置。是否 suspend 是个判断:只有当 checkpoint 已存在、且这块闲置 KV 会挡住别的等待工作时才 suspend,否则就留着。短等待倾向留(省 reload),长等待倾向 suspend(腾出的显存能服务别人)。suspend 只留轻量元数据。结果回来就重新入调度、有空间就 reload。",
      "Component three handles the interception itself, and it involves a real tradeoff. When a request makes an external call, its KV goes idle. We suspend it only when two things hold: a checkpoint already exists, and keeping those idle blocks on the GPU would actually block other waiting work. Otherwise we just retain it. The intuition: short waits favor retention to avoid reload overhead, while long waits favor suspension because the freed memory can serve other requests. Suspending records the external-call state and checkpoint generation, removes the request from the running set, and releases the GPU KV — leaving only lightweight metadata. When the result comes back, it re-enters the scheduler and reloads as soon as space is available.")

# ============================== SLIDE 11: GENERALITY ==============================
s = slide(); title(s, "One Substrate, Three Events")
bullets(s, [
    "The reload path is independent of WHY the GPU blocks were released — so the same mechanism generalizes:",
    "*External pause — tool / retrieval waits (the interception case above).",
    "*Reroute — a suspended request can resume on ANOTHER engine for load balancing; the receiver rebuilds the paged layout from the manifest.",
    "*Worker failure — if an engine loses its GPU KV, the request still resumes as long as the host checkpoint survives. A surviving engine reloads + replays the suffix — no full re-prefill.",
    "Same code path; only the trigger differs.",
], size=19)
notes(s,
      "通用性:reload 路径不关心 GPU 块为什么被释放,所以一套机制覆盖三种事件——外部暂停、跨引擎 reroute(负载均衡)、worker 故障恢复(host checkpoint 还在就能在别的引擎恢复,不用重算)。强调:同一条代码路径,只是触发原因不同。这是 Ferry 的「以一驭三」。",
      "Because the reload path doesn't care why the GPU blocks were released, one substrate covers three different serving events. First, external pauses — the interception case. Second, reroute: a suspended request can come back on a different engine for load balancing, and the receiver rebuilds the layout from the manifest. Third, worker failure: if an engine dies and loses its GPU KV, the request still resumes as long as the host checkpoint survives — a surviving engine reloads and replays the suffix, with no full re-prefill. It's literally the same code path; only the trigger differs. That generality is a big part of why we frame Ferry as a substrate rather than a single feature.")

# ============================== SLIDE 12: CODE STRUCTURE ==============================
s = slide(); title(s, "Implementation — What We Changed in vLLM", accent=CORAL)
codebox(s, [
    ("+  NEW — host checkpoint store", ""),
    ("   kv_checkpoint_pool.py", "758 LOC"),
    ("     save · restore · manifest · evict", ""),
    ("", ""),
    ("~  MODIFIED — hooks (~1.2K LOC)", ""),
    ("   gpu_model_runner.py", "~830 · per-step publish + reload"),
    ("   engine/core.py", "~190 · status + preempt-queue bus"),
    ("   sched/scheduler.py", "~170 · capacity-preempt -> reload"),
    ("", ""),
    ("-  UNCHANGED", ""),
    ("   model · attention kernels · paged layout", ""),
], title_txt="Diff vs. vanilla vLLM", height=4.95, width=7.4)
card(s, 8.3, 1.8, 4.35, 1.5, "STACK",
     "vLLM v1 fork · Qwen2.5-14B / 7B · A6000 fp16 · CUDA copy-stream + pinned host RAM · /dev/shm IPC (4 channels)")
card(s, 8.3, 3.45, 4.35, 1.35, "EVAL INFRA (outside the engine)",
     "router.py ~580 LOC — cross-engine reroute · least_load · + microbench harness", head_color=CORAL)
card(s, 8.3, 4.95, 4.35, 1.5, "TAKEAWAY",
     "~1.9K LOC of substrate in vLLM: one new subsystem + hooks in the runner / scheduler / engine. The model path is byte-for-byte vanilla vLLM — we change only WHEN KV leaves / returns to the GPU, not how the model runs.")
notes(s,
      "实现结构页(按「相对 vLLM 的 diff」组织,最清晰)。左边深色框分三层:➊ 新增——kv_checkpoint_pool.py(758 行,host checkpoint store:save/restore/manifest/eviction);➋ 改动的钩子(~1.2K 行)——gpu_model_runner.py(~830,把 publish 和 reload 接进每步执行循环,这是最大也最核心的一块)、engine/core.py(~190,engine 状态写出 + preempt 队列总线,撑 reroute/failover)、scheduler.py(~170,capacity-preempt 改成 reload 而非重算);➌ 没动的——模型 forward、attention kernel、paged layout、tokenizer。右边:技术栈、引擎之外的 eval infra(router 577 行)、takeaway。substrate 本体合计约 1.9K 行(另有探索性的 picker ~520 行没进贡献)。强调那句:模型执行路径和原版 vLLM 逐字节一样,我们只改 KV 何时搬。",
      "Here's what we changed in vLLM, as a diff. One new file: kv_checkpoint_pool.py, 758 lines — the host checkpoint store with save, restore, manifest, and eviction. Then the hooks, about twelve hundred lines. The biggest and most important is in gpu_model_runner.py — roughly 830 lines that wire the per-step publish and the reload into vLLM's execution loop. The engine core adds about 190 lines for writing engine status and draining the preempt queue, which supports reroute and failover. The scheduler adds about 170 to turn a capacity preemption into a reload instead of a recompute. Critically, the bottom group is unchanged: the model forward, the attention kernels, the paged-attention layout, the tokenizer. Outside the engine, the router is evaluation infrastructure. In total the substrate is about 1.9K lines. The point to land: the model execution path is byte-for-byte vanilla vLLM — we only change when KV leaves and returns to the GPU, not how the model runs.")

# ============================== SLIDE 13: DELTA + COPY STREAM ==============================
s = slide(); title(s, "Implementation — Delta Checkpoint + Overlapped Copy",
                    accent=CORAL)
codebox(s, [
    ("# only publish NEW blocks since last gen (delta)", ""),
    ("delta = block_ids[prev_n:]", ""),
    ("", ""),
    ("# copy stream waits on a decode-stream event,", ""),
    ("# decode never waits on copy -> they overlap", ""),
    ("ev = current_stream.record_event()", ""),
    ("with cuda.stream(copy_stream):", ""),
    ("    copy_stream.wait_event(ev)", ""),
    ("    for L in layers:", ""),
    ("        buf[L] = gpu_kv[L][:, blk_idx].clone()", ""),
    ("        pinned[L].copy_(buf[L], non_blocking=True)", ""),
], title_txt="kv_checkpoint_pool.py  (simplified)", height=4.6, width=7.4)
bullets(s, [
    "Delta: most steps copy ~1 new block, not the whole KV.",
    "Dedicated CUDA copy stream + an event on the decode stream:",
    (1, "decode N+1 overlaps the gather + PCIe copy."),
    "Per-layer gather → pinned host buffer → non-blocking D2H.",
    "*Result: +0.79 ms TPOT = 0.24% overhead (4 in-flight).",
], left=8.4, top=1.85, width=4.3, size=16)
notes(s,
      "实现细节一:delta + 重叠拷贝(对应 Q1 的低开销)。两个关键工程点:(1) delta——只发布上次之后的新块,稳态每步基本就 1 个块;(2) 用独立 copy stream + 在 decode stream 上记 event,让 copy 等 decode、decode 不等 copy,于是 decode N+1 和拷贝/PCIe 重叠。逐层 gather→pinned buffer→非阻塞 D2H。结果就是 0.24% 开销。代码是简化版,讲清楚机制即可。",
      "Two engineering details make checkpointing nearly free. First, delta checkpointing: we only publish blocks that are new since the last generation, so in steady state each step copies about one block, not the whole cache. Second, overlap: we run the copy on a dedicated CUDA stream and record an event on the decode stream. The copy stream waits on that event, but decode never waits on the copy — so the next decode step overlaps the gather and the PCIe transfer. We gather per layer into a pinned host buffer with a non-blocking device-to-host copy. The net effect is 0.79 milliseconds of added per-token latency under four concurrent requests — a 0.24 percent overhead.")

# ============================== SLIDE 14: RELOAD HOOK ==============================
s = slide(); title(s, "Implementation — Reload Path + Scheduler Hook",
                    accent=CORAL)
codebox(s, [
    ("# FT_CAPACITY_PREEMPT_RELOAD: KV exhausted &", ""),
    ("# waiting work blocked -> preempt a victim,", ""),
    ("# release its KV; it RELOADS later (not recompute)", ""),
    ("if kv_exhausted and waiting and reload_enabled:", ""),
    ("    victim = pick_capacity_victim(running)", ""),
    ("    release_kv(victim)        # checkpoint survives", ""),
    ("    mark_resumable(victim)    # reload, not re-prefill", ""),
], title_txt="sched/scheduler.py  (simplified)", height=3.4, width=7.6)
bullets(s, [
    "The hook turns preemption into a routine, cheap action.",
    "*Reload is 52–121× cheaper than re-prefill (next section), so releasing a victim's KV is no longer catastrophic.",
    "restore_checkpoint(): manifest → fresh blocks → H2D copy → rebuild paged layout → bounded replay.",
    "No new scheduling policy needed — a simple capacity-driven trigger suffices.",
], left=0.7, top=5.4, width=12.0, size=17)
notes(s,
      "实现细节二:reload 路径 + 调度器钩子。关键 flag FT_CAPACITY_PREEMPT_RELOAD:显存满且有等待工作时,踢一个 victim、释放它的 KV——但因为有 checkpoint,它之后是 reload 回来,不是重算。这把「抢占」变成便宜的常规操作(因为 reload 比 re-prefill 便宜 52-121×,下一节给数)。restore 流程:manifest→新块→H2D→重建 layout→有界 replay。强调:不需要新调度策略,简单的容量触发就够。",
      "The reload path plugs into the scheduler through one hook. When KV memory is exhausted and there's waiting work, we preempt a victim and release its KV — but because a checkpoint exists, that victim later reloads instead of recomputing. This is what turns preemption from a catastrophe into a routine, cheap scheduling action, because reload is one to two orders of magnitude cheaper than re-prefill, which I'll show next. The restore routine reads the manifest, allocates fresh blocks, copies K/V back, rebuilds the layout, and replays the short suffix. Importantly, we did not need a new scheduling policy — a simple capacity-driven trigger is enough to demonstrate the benefit.")

# ============================== SLIDE 15: EVAL METHODOLOGY ==============================
s = slide(); title(s, "Evaluation Methodology")
bullets(s, [
    "*Models:  Qwen2.5-14B (32K ctx, main end-to-end) · Qwen2.5-7B (up to 64K, reload microbench) · fp16.",
    "*Hardware:  2 × NVIDIA A6000 (48 GB) PCIe.",
    "*Workloads:",
    (1, "Microbench — synthetic 1K–64K prompts to isolate re-prefill vs reload."),
    (1, "Tool-pause — arxivsumm long prompts; decode → pause D seconds → resume."),
    (1, "Disruption — 2 engines, kill one mid-decode."),
    "*Baselines:  recompute (vanilla vLLM) · GPU prefix caching (APC) · reroute + re-prefill · (reactive swap, by construction).",
    "*Metrics:  TPOT overhead · reload-vs-reprefill · resume latency · goodput / P95 · failover gap.",
], size=17)
notes(s,
      "评测方法学(对应 5 分那项,简洁讲清)。模型:主实验 14B/32K,reload microbench 用 7B 上到 64K。硬件:2×A6000。三个 workload:合成 microbench、tool-pause(arxivsumm)、disruption(双引擎杀一个)。基线:recompute、APC、reroute+reprefill,外加 by-construction 的 reactive swap。指标:TPOT 开销、reload vs reprefill、resume latency、goodput/P95、failover gap。",
      "Quickly, the methodology. We use Qwen2.5-14B at 32K context for the end-to-end experiments, and 7B up to 64K for the reload microbenchmark, all in fp16, on two A6000s. Three workloads: a synthetic microbenchmark to isolate reload versus re-prefill; a tool-pause workload built on long arxiv-summarization prompts where each request decodes, pauses for a controlled duration, then resumes; and a disruption workload that kills one of two engines mid-decode. Baselines are recompute, GPU prefix caching, and reroute-with-re-prefill, plus reactive swap analyzed by construction. Metrics are decode overhead, reload-versus-reprefill, resume latency, goodput and P95 latency, and the failover gap.")

# ============================== SLIDE 16: MICROBENCH TABLES ==============================
s = slide(); title(s, "Results 1 — Mechanism Microbenchmarks")
# Table 1
t1 = s.shapes.add_table(3, 3, Inches(0.7), Inches(1.8), Inches(6.0),
                        Inches(1.5)).table
for j, h in enumerate(["Workload", "vLLM → +ckpt", "Δ TPOT"]):
    c = t1.cell(0, j); c.text = h
    c.text_frame.paragraphs[0].runs[0].font.bold = True
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    c.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
    c.fill.solid(); c.fill.fore_color.rgb = BLUE
for i, row in enumerate([("Sequential (1)", "46.07 → 46.07 ms", "−0.00 ms"),
                         ("Concurrent (4)", "334.32 → 335.11 ms",
                          "+0.79 ms (0.24%)")], start=1):
    for j, v in enumerate(row):
        c = t1.cell(i, j); c.text = v
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        c.fill.solid(); c.fill.fore_color.rgb = WHITE
# Table 2
t2 = s.shapes.add_table(5, 3, Inches(7.1), Inches(1.8), Inches(5.5),
                        Inches(3.0)).table
for j, h in enumerate(["Context", "Re-prefill / Reload", "Speedup"]):
    c = t2.cell(0, j); c.text = h
    c.text_frame.paragraphs[0].runs[0].font.bold = True
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    c.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
    c.fill.solid(); c.fill.fore_color.rgb = CORAL
for i, row in enumerate([("1K", "0.49 s / 9.4 ms", "52×"),
                         ("16K", "10.40 s / 153.7 ms", "68×"),
                         ("32K", "26.31 s / 306.4 ms", "86×"),
                         ("64K", "73.83 s / 611.8 ms", "121×")], start=1):
    for j, v in enumerate(row):
        c = t2.cell(i, j); c.text = v
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        c.fill.solid(); c.fill.fore_color.rgb = WHITE
bullets(s, [
    "*Table 1: continuous checkpointing adds at most 0.24% TPOT — essentially free.",
    "*Table 2: reload is 52–121× cheaper than re-prefill; the gap widens with context.",
    "→ The mechanism is cheap enough to make suspension & recovery routine.",
], top=5.0, size=17)
notes(s,
      "结果一:机制 microbench(回答 Q1)。左表:连续 checkpoint 开销——单请求 0,4 并发 +0.79ms=0.24%,基本免费。右表:reload vs re-prefill,从 1K 的 52× 到 64K 的 121×,上下文越长差距越大。结论:机制足够便宜,所以 suspend 和恢复都能当常规操作。",
      "First results: the mechanism in isolation. The left table is the checkpointing overhead — zero for a single request, and just 0.79 milliseconds, or 0.24 percent, with four concurrent requests. Essentially free. The right table compares reload against re-prefill across context lengths: reload is 52 times cheaper at 1K, rising to 121 times cheaper at 64K — the gap widens as context grows. Together these say the mechanism is cheap enough that we can treat suspension and recovery as routine operations rather than expensive exceptions.")

# ============================== SLIDE 17: TOOL-PAUSE ==============================
s = slide(); title(s, "Results 2 — Tool-Pause Resume Latency")
picture(s, "fig_icept_crossover.png", 0.7, 1.7, height=4.3)
bullets(s, [
    "*Ferry: flat-low, ~2.5–3 s at every pause duration.",
    "APC: starts ~5 s, degrades to 8–9 s once the idle KV is evicted.",
    "Recompute: 10–15 s throughout (re-prefills every resume).",
    "*Ferry is lowest at every pause; the gap grows with pause length — up to ~5× vs recompute.",
    "Host checkpoint survives GPU pressure, so resume cost is pause-independent.",
], left=7.6, top=1.8, width=5.2, size=16)
notes(s,
      "结果二:tool-pause 的 resume 延迟(回答 Q2 的延迟侧)。蓝线 Ferry 不管暂停多久都稳在 ~2.5-3s;APC 短暂停 ~5s,一长被挤掉就涨到 8-9s;recompute 一直 10-15s。Ferry 在每个点都最低,暂停越长优势越大,对 recompute 最多 ~5×。原因:host checkpoint 不受显存压力影响,所以恢复代价与暂停时长无关。",
      "Second, resume latency on the tool-pause workload. The blue line is Ferry: flat and low, around two-and-a-half to three seconds regardless of how long the request paused. Prefix caching starts around five seconds but degrades to eight or nine once the idle KV gets evicted under memory pressure. Recompute stays at ten to fifteen seconds because it re-prefills on every resume. Ferry is the lowest at every pause duration, and the advantage grows with pause length — up to about five times faster than recompute. The reason is that the host checkpoint is immune to GPU memory pressure, so Ferry's resume cost is essentially independent of the pause.")

# ============================== SLIDE 18: THROUGHPUT ==============================
s = slide(); title(s, "Results 2 — Throughput & Tail Latency")
picture(s, "fig_icept_qps.png", 0.7, 1.8, height=3.6)
bullets(s, [
    "*Near saturation, Ferry sustains 1.17–1.23× the completed throughput of recompute.",
    "Lower P95 work latency at high arrival rates.",
    "Avoiding re-prefill on resume frees engine capacity the baseline must spend recomputing.",
], left=0.7, top=5.6, width=12.0, size=17)
notes(s,
      "结果二补充:吞吐和尾延迟。近饱和时 Ferry 比 recompute 高 1.17-1.23× 的完成吞吐,P95 work latency 也更低。本质原因:不在恢复时重算 prefill,就把算力省下来服务更多请求,负载越高差距越明显。",
      "Still on the interception workload, this is throughput and tail latency versus arrival rate. Near saturation, Ferry sustains about 1.17 to 1.23 times the completed throughput of the recompute baseline, and it keeps P95 work latency lower at high load. The mechanism is the same: by not re-prefilling on every resume, Ferry frees engine capacity that the baseline has to spend on recomputation, and that gap widens as the system approaches saturation.")

# ============================== SLIDE 19: DISRUPTION ==============================
s = slide(); title(s, "Results 3 — Disruption Recovery")
picture(s, "fig_failover_bar.png", 0.7, 1.7, height=4.4)
bullets(s, [
    "Kill 1 of 2 engines mid-decode; measure the failover gap to the next token on a surviving engine.",
    "Detection latency (~0.7 s) is common to both systems.",
    "*Reroute + re-prefill: 8.9 s resume.   Ferry reload: 5.1 s resume → 1.7× faster on the resume path.",
    "*Total failover gap: 9.6 s → 5.7 s.",
    "Host checkpoint survives the dead GPU — recovery never touches the failed engine.",
], left=7.2, top=1.8, width=5.6, size=16)
notes(s,
      "结果三:故障恢复(回答 Q3)。双引擎杀一个,量 failover gap(从被杀到在另一台产出下一个 token)。检测延迟 ~0.7s 两者相同。reroute+reprefill 的 resume 8.9s,Ferry reload 5.1s,resume 路径上快 1.7×;总 gap 从 9.6s 降到 5.7s。关键:host checkpoint 在 GPU 死后还在,恢复完全不依赖坏掉的引擎。注意这里用的是修正后的数(1.7×、5.1s),跟图一致。",
      "Third, disruption recovery — this answers our generality question. We kill one of two engines mid-decode and measure the failover gap, the time until the affected request produces its next token on a surviving engine. About 0.7 seconds of detection latency is common to both systems. The reroute-and-reprefill baseline pays 8.9 seconds of resume cost; Ferry reloads from the host checkpoint in 5.1 seconds — 1.7 times faster on the resume path. The total failover gap drops from 9.6 to 5.7 seconds. The key point is that the host checkpoint survives the dead GPU, so recovery never needs to touch the failed engine — the same reload path we built for pauses just works for failures.")

# ============================== SLIDE 20: DEMO ==============================
s = slide(); title(s, "Demo — Tool-Pause Resume: Ferry vs Recompute", accent=CORAL)
box_with_text(s, 0.7, 1.65, 12.0, 0.8,
              ["▶  screen recording plays here  —  demo_toolpause.sh"],
              fill=RGBColor(0xFF, 0xF3, 0xE0), line_color=CORAL, size=15)
bullets(s, [
    "Setup: one engine, a 6–12K-token request decodes, issues a 15 s external pause, then resumes.",
    "Three panes on screen: resume latency + reload validity · /dev/shm checkpoint store growing · nvidia-smi (GPU freed on suspend).",
    "*Ferry: 8/8 requests reload from the host checkpoint — resume ≈ 2–3 s.",
    (1, "engine log: \"reload took 1 step · 6176 tokens restored\""),
    "*Recompute: re-prefills the whole long prompt every resume — ≈ 10–15 s.",
    "Same workload; only the resume strategy differs.",
], top=2.75, size=17)
notes(s,
      "Demo 页(方案 A,录屏)。这页放 tool-pause 录屏,讲稿配着放。设置:单引擎、一个 6-12K 长请求、解码一会儿→发起 15s 外部暂停→恢复。屏幕三个窗口:主窗口跑 demo_toolpause.sh(打印 resume latency + reload validity);窗口2 看 /dev/shm/vllm_ft_checkpoints 块数在涨;窗口3 看 nvidia-smi 显存在 suspend 时掉下去。重点念出来的两行:Ferry「8/8 reloaded / 0 fell back, resume≈2-3s」、引擎日志「reload took 1 step, 6176 tokens restored」;对照 recompute resume≈10-15s。脚本在 experiments_v2/CSE 232/demo_toolpause.sh。注意:脚本开头会清场杀掉旧引擎,GPU 别的任务在跑时别启动。",
      "Let me show this live. The setup is one engine serving a 6-to-12-thousand-token request. It decodes for a moment, then issues a 15-second external pause — simulating a tool call — and then resumes. Watch three things on screen: the main terminal printing resume latency and reload validity, the shared-memory checkpoint store on the left growing as completed blocks are published, and nvidia-smi on the right showing GPU memory dropping the instant the request suspends. With Ferry, all eight intercepted requests reload from the host checkpoint — you can see in the engine log, 'reload took one step, 6176 tokens restored' — and the resume completes in about two to three seconds. Now the same workload with vanilla vLLM: on resume it re-prefills the entire long prompt, so it takes ten to fifteen seconds. Same workload, same hardware — the only difference is whether we reload a checkpoint or recompute from scratch.")

# ============================== SLIDE 21: DISCUSSION ==============================
s = slide(); title(s, "Discussion & Limitations")
bullets(s, [
    "*No new scheduling policy needed. A simple capacity-driven trigger suffices; SLO-aware scheduling is orthogonal and can sit on the same reload path.",
    "Limitations:",
    (1, "Two engines, one hardware platform, one model size — larger clusters need coordinating concurrent reroutes."),
    (1, "Assumes the host checkpoint store survives a failure; full fault tolerance needs replication (out of scope)."),
    (1, "Assumes PCIe is not saturated during decode — may not hold under prefill/decode disaggregation or CPU-offload, where checkpoint traffic must be co-scheduled."),
], size=18)
notes(s,
      "讨论与局限,诚实交代。Ferry 不需要新调度策略,简单容量触发就能体现收益,SLO-aware 调度是正交的、可叠在同一 reload 路径上。三个局限:规模(2 引擎/1 平台/1 模型)、假设 checkpoint store 不丢(否则要副本)、假设 decode 时 PCIe 不饱和(P/D 分离或 CPU offload 下不成立,需要把 checkpoint 流量一起调度)。",
      "A few points of discussion. First, Ferry needs no new scheduling policy — a simple capacity-driven trigger already shows the benefit, and SLO-aware scheduling is orthogonal; it can be layered on the same reload path. As for limitations: our evaluation uses two engines, one platform, and one model size, so scaling out would require coordinating concurrent reroutes across many peers. We assume the host checkpoint store survives a failure — full fault tolerance would need replication, which is out of scope. And we assume PCIe is not saturated during decode; that holds in our setup but may not under prefill-decode disaggregation or CPU offload, where checkpoint traffic would have to be co-scheduled with other transfers.")

# ============================== SLIDE 22: CONCLUSION ==============================
s = slide(); title(s, "Conclusion")
bullets(s, [
    "*Problem: a KV-residency dilemma in long-context augmented serving — retain wastes HBM, discard forces re-prefill.",
    "*Ferry: proactive, block-granular checkpointing decouples KV lifetime from GPU residency.",
    (1, "Suspend releases GPU KV immediately; resume = reload + bounded replay; one substrate covers interception, reroute, and worker failure."),
    "*Results:",
    (1, "Up to 5× lower resume latency · 1.17–1.23× throughput · 1.7× faster failover recovery · only 0.24% steady-state overhead."),
    "Effect comes from WHEN KV moves — off the critical path — not from changing the model.",
], size=18)
caption(s, "Thank you — questions?", top=6.7, size=16, color=BLUE)
notes(s,
      "结论页。回收三点:问题(KV 驻留两难)、做法(主动按块 checkpoint,解耦生命周期与 GPU 驻留;suspend 即放、resume=reload+有界 replay;一套底座覆盖拦截/reroute/故障)、结果(resume 最多 5×、吞吐 1.17-1.23×、故障恢复 1.7×、开销 0.24%)。最后一句点睛:收益来自「KV 什么时候搬」,不是改模型。然后进 Q&A。",
      "To conclude. The problem is a KV-residency dilemma in long-context augmented serving: retaining wastes scarce memory, discarding forces a full re-prefill. Ferry solves it with proactive, block-granular checkpointing that decouples KV lifetime from GPU residency. Suspension releases GPU KV immediately, resume is a reload plus a bounded replay, and one substrate covers interception, reroute, and worker failure. In numbers: up to five times lower resume latency, up to 1.23 times the throughput, 1.7 times faster failover recovery, and only 0.24 percent steady-state overhead. The whole effect comes from changing when KV moves — taking it off the critical path — not from changing the model. Thank you; I'm happy to take questions.")

out = os.path.join(HERE, "Ferry_CSE232_Slides.pptx")
prs.save(out)
print("saved:", out, "| slides:", len(prs.slides._sldIdLst))
